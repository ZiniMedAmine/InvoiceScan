"""
Conversion of the extracted JSON data into downloadable files.

Every exporter takes the document type and the parsed JSON data and returns
the content of the file as bytes, so files can be zipped or stored directly
without touching the disk.
"""

import csv
import io
import json
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from typing import Any

from docx import Document


def flatten(data: Any, prefix: str = "") -> Iterator[tuple[str, Any]]:
    """
    Yield ``(dotted.key, value)`` pairs for every leaf of a nested structure.

    >>> dict(flatten({"client": {"name": "ACME"}, "items": ["a", "b"]}))
    {'client.name': 'ACME', 'items.0': 'a', 'items.1': 'b'}
    """
    if isinstance(data, dict):
        items = data.items()
    elif isinstance(data, list):
        items = enumerate(data)
    else:
        yield prefix, data
        return

    for key, value in items:
        yield from flatten(value, f"{prefix}.{key}" if prefix else str(key))


def to_json(document_type: str, data: dict) -> bytes:
    return json.dumps(data, indent=2, ensure_ascii=False).encode("utf-8")


def to_csv(document_type: str, data: dict) -> bytes:
    """Two-column CSV (``field``, ``value``) with one row per leaf value."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["field", "value"])
    writer.writerow(["document_type", document_type])
    writer.writerows(flatten(data))
    # BOM so that Excel detects UTF-8 and displays accented characters.
    return buffer.getvalue().encode("utf-8-sig")


def to_docx(document_type: str, data: dict) -> bytes:
    """Word document with one heading per section of the JSON data."""
    document = Document()
    document.add_heading(document_type or "Document", level=0)
    _add_docx_section(document, data, level=1)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_docx_section(document, data: Any, level: int) -> None:
    items = data.items() if isinstance(data, dict) else enumerate(data, start=1)
    for key, value in items:
        if isinstance(value, (dict, list)):
            document.add_heading(str(key), level=min(level, 9))
            _add_docx_section(document, value, level + 1)
        else:
            paragraph = document.add_paragraph()
            paragraph.add_run(f"{key}: ").bold = True
            paragraph.add_run(str(value))


@dataclass(frozen=True)
class ExportFormat:
    extension: str
    render: Callable[[str, dict], bytes]


# Keys match ``ExportedFile.Format`` and the checkboxes of the review page.
EXPORT_FORMATS: dict[str, ExportFormat] = {
    "json": ExportFormat("json", to_json),
    "csv": ExportFormat("csv", to_csv),
    "word": ExportFormat("docx", to_docx),
}
